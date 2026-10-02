"""Documento "Encaminhamentos do SAMU" — uma página por hospital de destino.

Fonte: relatório 115 do vReport do vSky ("Atendimentos por Hospital de
Destino"), exportado em Excel com os filtros:

    Cliente: SAMU ESPIRITO SANTO · Considerar hospitais sem atendimento: Não
    Região: METROPOLITANO 1, 2 e 3 · Município e Hospital: vazios

A planilha repete o cabeçalho a cada página do relatório; cada linha é um
hospital com contagens por cor, origem e tipo. O documento usa:

    Total               pacientes que o hospital recebeu do SAMU
    Inter-hospitalar    atendimento secundário (transferências de PA e hospitais)
    Pré-hospitalar      atendimento primário (APH)

Funções puras: ler a planilha, escrever números por extenso, nomear o
hospital e montar o .docx.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date

COLUNAS = {"hospital": "Hospital", "municipio": "Município", "total": "Total",
           "pre": "Pré-hospitalar", "inter": "Inter-hospitalar"}

FILTROS_VSKY = (
    ("Cliente", "SAMU ESPIRITO SANTO"),
    ("Considerar Hospitais sem atendimento", "Não"),
    ("Região", "METROPOLITANO 1, METROPOLITANO 2, METROPOLITANO 3"),
    ("Município", "vazio"),
    ("Hospital", "vazio"),
)
URL_RELATORIO = ("https://gestao-es.vskysamu.com.br/vreport-management-web/"
                 "#/formularioRelatorio/115")

COMPROMISSO = ("Atendimento aos pacientes dos Prontos atendimentos de referência do "
               "Hospital, pactuado no Comitê Gestor de Urgência e Emergência e do SAMU")
META = "100% dos pacientes acolhidos"
PONTOS = "10"
INSTRUMENTO = "Relatório do NERUE"


# ------------------------------------------------------------------ planilha

def _chave(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _inteiro(valor) -> int:
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return 0


def ler_relatorio(conteudo: bytes) -> list[dict]:
    """Linhas de hospital do .xls/.xlsx exportado pelo vReport.

    Localiza as colunas pelo cabeçalho (que se repete a cada página) em vez
    de posição fixa, e ignora títulos, linhas em branco e o rodapé de total.
    """
    import pandas as pd

    try:
        bruto = pd.read_excel(io.BytesIO(conteudo), header=None, sheet_name=None)
    except Exception as exc:  # noqa: BLE001 — arquivo que não é planilha
        raise ValueError("Não foi possível ler a planilha. Envie o Report.xls exportado "
                         "do vReport (Save → Microsoft Excel).") from exc
    linhas: list[dict] = []
    for df in bruto.values():
        mapa: dict[str, int] | None = None
        for _, row in df.iterrows():
            valores = list(row.values)
            chaves = [_chave(v) for v in valores]
            if "hospital" in chaves and "total" in chaves:
                mapa = {campo: chaves.index(_chave(rotulo))
                        for campo, rotulo in COLUNAS.items() if _chave(rotulo) in chaves}
                continue
            if mapa is None or "hospital" not in mapa:
                continue
            nome = valores[mapa["hospital"]]
            if not isinstance(nome, str) or not nome.strip():
                continue
            linhas.append({
                "hospital": re.sub(r"\s+", " ", nome).strip(),
                "municipio": str(valores[mapa["municipio"]]).strip()
                if "municipio" in mapa and isinstance(valores[mapa["municipio"]], str) else "",
                "total": _inteiro(valores[mapa["total"]]) if "total" in mapa else 0,
                "pre": _inteiro(valores[mapa["pre"]]) if "pre" in mapa else 0,
                "inter": _inteiro(valores[mapa["inter"]]) if "inter" in mapa else 0,
            })
    if not linhas:
        raise ValueError("A planilha não tem o formato do relatório "
                         "\"Atendimentos por Hospital de Destino\".")
    return linhas


# ------------------------------------------------------------------ extenso

_UNIDADES = ("zero", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito",
             "nove", "dez", "onze", "doze", "treze", "catorze", "quinze", "dezesseis",
             "dezessete", "dezoito", "dezenove")
_DEZENAS = ("", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta",
            "oitenta", "noventa")
_CENTENAS = ("", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos",
             "seiscentos", "setecentos", "oitocentos", "novecentos")


def _ate_mil(n: int) -> str:
    if n < 20:
        return _UNIDADES[n]
    if n < 100:
        d, u = divmod(n, 10)
        return _DEZENAS[d] + (f" e {_UNIDADES[u]}" if u else "")
    if n == 100:
        return "cem"
    c, resto = divmod(n, 100)
    return _CENTENAS[c] + (f" e {_ate_mil(resto)}" if resto else "")


def extenso(n: int) -> str:
    """Número inteiro por extenso (0 a 999.999), como em documento oficial."""
    if n < 0:
        return "menos " + extenso(-n)
    if n < 1000:
        return _ate_mil(n)
    milhar, resto = divmod(n, 1000)
    prefixo = "mil" if milhar == 1 else f"{_ate_mil(milhar)} mil"
    if not resto:
        return prefixo
    # "e" só quando o resto é uma centena redonda ou menor que 100
    liga = " e " if resto < 100 or resto % 100 == 0 else " "
    return prefixo + liga + _ate_mil(resto)


def quantidade(n: int, singular: str = "paciente", plural: str = "pacientes") -> str:
    return f"{n} ({extenso(n)}) {singular if n == 1 else plural}"


# ------------------------------------------------------------------ nomes

_PEQUENAS = {"de", "da", "do", "das", "dos", "e"}
_ACENTOS = {
    "urgencia": "Urgência", "emergencia": "Emergência", "atencao": "Atenção",
    "clinica": "Clínica", "clinicas": "Clínicas", "evangelico": "Evangélico",
    "antonio": "Antônio", "gloria": "Glória", "sao": "São", "conceicao": "Conceição",
    "policia": "Polícia", "associacao": "Associação", "vitoria": "Vitória",
    "jetiba": "Jetibá", "piuma": "Piúma", "familia": "Família", "itaguacu": "Itaguaçu",
    "cassia": "Cássia", "saude": "Saúde", "joao": "João", "presidio": "Presídio",
    "municipio": "Município", "regiao": "Região",
}


def _palavra(p: str, primeira: bool) -> str:
    base = p.lower()
    if base in _PEQUENAS and not primeira:
        return base
    sem = _chave(base).strip(".")
    if sem in _ACENTOS:
        return _ACENTOS[sem] + ("." if base.endswith(".") else "")
    return base[:1].upper() + base[1:]


def separar_sigla(nome_vsky: str) -> tuple[str, str | None]:
    """'HOSPITAL ESTADUAL CENTRAL - HEC' → ('HOSPITAL ESTADUAL CENTRAL', 'HEC');
    'HOSPITAL … CONCEICAO - PIUMA' → (nome inteiro, None)."""
    partes = nome_vsky.rsplit(" - ", 1)
    sigla = partes[1].strip() if len(partes) == 2 else ""
    # sigla começa pela inicial do nome (HEC de Hospital…); "- PIUMA" é cidade
    if re.fullmatch(r"[A-ZÀ-Ú.]{2,10}", sigla) and sigla[0] == partes[0].strip()[:1]:
        return partes[0].strip(), sigla
    return nome_vsky.strip(), None


def nome_documento(nome_vsky: str) -> tuple[str, str | None]:
    """Nome em caixa normal para o título e a sigla, se houver."""
    nome, sigla = separar_sigla(nome_vsky)
    palavras = nome.split()
    return " ".join(_palavra(p, i == 0) for i, p in enumerate(palavras)), sigla


def incluir_por_padrao(nome_vsky: str) -> bool:
    """Sugestão inicial de quem entra no documento: hospitais estaduais e os
    que já aparecem no modelo. A seleção fica salva e é ajustada na tela."""
    n = _chave(nome_vsky)
    return n.startswith("hospital estadual") or n.endswith("- hevv")


def titulo(nome: str, sigla: str | None) -> str:
    return f"{nome} ({sigla})" if sigla else nome


def _artigo(nome: str) -> str:
    primeira = _chave(nome).split(" ")[0] if nome else ""
    return "a" if primeira in {"maternidade", "upa", "unidade", "policlinica"} else "o"


def memoria(nome: str, sigla: str | None, total: int, pre: int, inter: int) -> list[tuple]:
    """Texto da memória de cálculo em trechos (texto, negrito, sublinhado)."""
    sujeito = sigla or nome
    art = _artigo(nome)
    return [
        ("Memória de Cálculo:", True, False),
        (" Nesse período, segundo dados registrados e coletados do sistema VELP "
         f"VSKYSAMU, {art} ", False, False),
        (f"{sujeito} recebeu {quantidade(total)} pelo SAMU 192", True, True),
        (f". Do total de pacientes que {art} {sujeito} recebeu no referido mês, ",
         False, False),
        (quantidade(inter), True, False),
        (" foram provenientes de Atendimento secundário (Transferências de Pronto "
         "Atendimentos e Hospitais); e ", False, False),
        (f"{pre} ({extenso(pre)})", True, False),
        (" provenientes de Atendimento primário (APH).", False, False),
    ]


# ------------------------------------------------------------------ documento

def gerar_docx(paginas: list[dict], periodo: str, logo: bytes | None = None) -> bytes:
    """Monta o .docx: uma página por hospital.

    Cada item de `paginas`: nome, sigla, total, pre, inter.
    `periodo` (ex.: 'setembro/2026') vai no rodapé de cada página.
    """
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    secao = doc.sections[0]
    secao.orientation = WD_ORIENT.PORTRAIT
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    for lado in ("left_margin", "right_margin"):
        setattr(secao, lado, Cm(2))
    secao.top_margin, secao.bottom_margin = Cm(1.5), Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    # cabeçalho igual em todas as páginas: brasão + governo/secretaria
    cab = secao.header
    p = cab.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo:
        try:
            p.add_run().add_picture(io.BytesIO(logo), height=Cm(1.2))
            p = cab.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:  # noqa: BLE001 — imagem inválida: segue só com o texto
            pass
    for linha in ("GOVERNO DO ESTADO DO ESPÍRITO SANTO", "SECRETARIA DE ESTADO DA SAÚDE"):
        r = p.add_run(linha)
        r.font.size = Pt(6.5)
        if linha.startswith("GOVERNO"):
            r.add_break()

    rod = secao.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = rod.add_run(f"SAMU 192 ES — Encaminhamentos do SAMU — {periodo}")
    r.font.size = Pt(8)

    def borda(celula):
        tc = celula._tc.get_or_add_tcPr()
        bordas = OxmlElement("w:tcBorders")
        for lado in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{lado}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:color"), "000000")
            bordas.append(el)
        tc.append(bordas)

    larguras = [Cm(1.0), Cm(5.0), Cm(2.6), Cm(1.8), Cm(3.8), Cm(2.8)]
    for n, pg in enumerate(paginas):
        if n:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        t = doc.add_paragraph()
        t.paragraph_format.space_before = Pt(6)
        rt = t.add_run(titulo(pg["nome"], pg.get("sigla")))
        rt.bold = True
        rt.font.size = Pt(12)

        tabela = doc.add_table(rows=2, cols=6)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabela.autofit = False
        for i, coluna in enumerate(tabela.columns):     # Word e LibreOffice leem o grid
            coluna.width = larguras[i]
        cab_linha, dados = tabela.rows
        cab_linha.cells[0].merge(cab_linha.cells[1])
        titulos = ["COMPROMISSO", None, "META", "PONTOS", "INSTRUMENTO", "AVALIAÇÃO"]
        valores = ["01", COMPROMISSO, META, PONTOS, INSTRUMENTO, ""]
        for i in range(6):
            for linha in (cab_linha, dados):
                linha.cells[i].width = larguras[i]
                borda(linha.cells[i])
            if titulos[i]:
                pc = cab_linha.cells[i].paragraphs[0]
                pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rc = pc.add_run(titulos[i])
                rc.bold = True
                rc.font.size = Pt(10)
            for celula in (cab_linha.cells[i], dados.cells[i]):
                fmt = celula.paragraphs[0].paragraph_format
                fmt.space_after, fmt.line_spacing = Pt(2), 1.0
            pd_ = dados.cells[i].paragraphs[0]
            if i in (0, 3):
                pd_.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rd = pd_.add_run(valores[i])
            rd.font.size = Pt(10)

        obs = doc.add_paragraph()
        ro = obs.add_run("OBS: A pontuação conferida na “avaliação” deverá conter no "
                         "máximo uma casa decimal.")
        ro.font.size = Pt(10)

        corpo = doc.add_paragraph()
        corpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        corpo.paragraph_format.space_before = Pt(10)
        for texto, negrito, sublinhado in memoria(pg["nome"], pg.get("sigla"),
                                                  pg["total"], pg["pre"], pg["inter"]):
            rr = corpo.add_run(texto)
            rr.bold, rr.underline = negrito, sublinhado

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def periodo_dos_dados(competencia: date) -> tuple[date, date]:
    """Envio de outubro → dados de 01/09 a 30/09."""
    from app.modules.sesa.prazos import periodo_competencia_anterior

    return periodo_competencia_anterior(competencia)
