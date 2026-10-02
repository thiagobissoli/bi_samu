"""Documento "Saída de USA" — produção de Unidades de Suporte Avançado.

Item das prestações de contas dos convênios (007 e 008). É calculado dos
próprios registros do vSky que o sistema já importa, sem planilha nova.

A definição de **saída** é a mesma do módulo de Indicadores
(`tema_saidas_ambulancia`): um empenho que iniciou deslocamento — tem
unidade identificada, data da ocorrência e início de deslocamento. Daí
separa-se USA de USB pela coluna `recurso`, derivada do nome da viatura.

Funções puras: recebem o DataFrame do núcleo já recortado no período e
devolvem os totais e as quebras; a montagem do .xlsx não toca em banco.
"""

from __future__ import annotations

import io
from datetime import date

ORDEM_COR = ("vermelho", "amarelo", "verde", "nao_urgente", "orientacao_medica")
ROTULO_COR = {"vermelho": "Vermelho", "amarelo": "Amarelo", "verde": "Verde",
              "nao_urgente": "Não urgente", "orientacao_medica": "Orientação médica"}


def saidas(df):
    """Empenhos que saíram (iniciaram deslocamento), com unidade identificada."""
    return df[df["dt_ocorr"].notna()
              & df["unidade_curta"].notna() & df["unidade_curta"].ne("")
              & df["dt_inicio_deslocamento"].notna()]


def _contagem(serie) -> list[tuple[str, int]]:
    vc = serie.value_counts()
    return [(str(k), int(v)) for k, v in vc.items()]


def calcular(df, inicio: date, fim: date) -> dict:
    """Totais de saída no período (inicio..fim, pela data da ocorrência).

    Devolve total geral, USA, USB e, para as USA, as quebras por unidade,
    por município, por código (cor) e por dia.
    """
    base = saidas(df)
    if "dia" in base.columns:
        base = base[(base["dia"] >= inicio) & (base["dia"] <= fim)]
    usa = base[base["recurso"] == "USA"]
    usb = base[base["recurso"] == "USB"]
    dias = sorted({d for d in base["dia"].dropna().unique()})

    por_cor = []
    cor = usa["codigo_cor"] if "codigo_cor" in usa.columns else None
    if cor is not None:
        contagem = {str(k): int(v) for k, v in cor.value_counts().items()}
        for chave in ORDEM_COR:
            if contagem.get(chave):
                por_cor.append((ROTULO_COR.get(chave, chave), contagem[chave]))
        conhecidas = set(ROTULO_COR)
        for chave, n in contagem.items():
            if chave not in conhecidas:
                por_cor.append((chave, n))

    return {
        "inicio": inicio, "fim": fim,
        "total": len(base), "usa": len(usa), "usb": len(usb),
        "dias": len(dias),
        "por_unidade": _contagem(usa["unidade_curta"]),
        "por_municipio": _contagem(usa["cidade"].dropna()) if "cidade" in usa.columns else [],
        "por_cor": por_cor,
        "por_dia": [(d, int((usa["dia"] == d).sum())) for d in dias],
    }


def _pct(parte: int, total: int) -> float:
    return round(parte / total * 100, 1) if total else 0.0


def gerar_xlsx(dados: dict, periodo_rotulo: str, orgao: str = "SAMU 192 ES") -> bytes:
    """Planilha da Saída de USA: resumo e quebras, uma aba cada."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    titulo_fill = PatternFill("solid", fgColor="1F4E79")
    titulo_font = Font(bold=True, color="FFFFFF", size=12)
    cab_fill = PatternFill("solid", fgColor="D9E1F2")
    negrito = Font(bold=True)
    centro = Alignment(horizontal="center")

    wb = Workbook()

    def _aba(nome, titulo, colunas, linhas, total_label=None, total_valor=None):
        ws = wb.create_sheet(nome)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(colunas)))
        c = ws.cell(1, 1, titulo)
        c.fill, c.font = titulo_fill, titulo_font
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 22
        for j, nome_col in enumerate(colunas, start=1):
            cel = ws.cell(3, j, nome_col)
            cel.fill, cel.font = cab_fill, negrito
            if j > 1:
                cel.alignment = centro
        linha = 4
        for valores in linhas:
            for j, v in enumerate(valores, start=1):
                cel = ws.cell(linha, j, v)
                if j > 1:
                    cel.alignment = centro
            linha += 1
        if total_label is not None:
            cel = ws.cell(linha, 1, total_label)
            cel.font = negrito
            tv = ws.cell(linha, 2, total_valor)
            tv.font, tv.alignment = negrito, centro
        larguras = [max(len(str(colunas[j])), *(len(str(v[j])) for v in linhas)) + 2
                    if linhas else len(str(colunas[j])) + 2 for j in range(len(colunas))]
        for j, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(j)].width = max(10, min(45, w))
        return ws

    # Resumo
    ws = wb.active
    ws.title = "Resumo"
    ws.merge_cells("A1:C1")
    c = ws.cell(1, 1, f"Saída de USA — {orgao}")
    c.fill, c.font = titulo_fill, titulo_font
    ws.cell(2, 1, f"Período: {periodo_rotulo}").font = Font(italic=True)
    linhas = [
        ("Total de saídas de unidade", dados["total"], ""),
        ("Saídas de USA", dados["usa"], f'{_pct(dados["usa"], dados["total"])}%'),
        ("Saídas de USB", dados["usb"], f'{_pct(dados["usb"], dados["total"])}%'),
        ("Dias com registro", dados["dias"], ""),
        ("Média de saídas de USA por dia",
         round(dados["usa"] / dados["dias"], 1) if dados["dias"] else 0, ""),
    ]
    for j, nome_col in enumerate(("Indicador", "Valor", "%"), start=1):
        cel = ws.cell(4, j, nome_col)
        cel.fill, cel.font = cab_fill, negrito
        if j > 1:
            cel.alignment = centro
    for i, (rot, val, pct) in enumerate(linhas, start=5):
        ws.cell(i, 1, rot)
        ws.cell(i, 2, val).alignment = centro
        ws.cell(i, 3, pct).alignment = centro
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10

    total_usa = dados["usa"]
    _aba("Por unidade", "Saídas de USA por unidade", ("Unidade", "Saídas", "%"),
         [(u, n, f"{_pct(n, total_usa)}%") for u, n in dados["por_unidade"]],
         "Total", total_usa)
    _aba("Por município", "Saídas de USA por município", ("Município", "Saídas", "%"),
         [(m, n, f"{_pct(n, total_usa)}%") for m, n in dados["por_municipio"]],
         "Total", total_usa)
    _aba("Por código", "Saídas de USA por código", ("Código", "Saídas", "%"),
         [(c, n, f"{_pct(n, total_usa)}%") for c, n in dados["por_cor"]],
         "Total", total_usa)
    _aba("Por dia", "Saídas de USA por dia", ("Dia", "Saídas"),
         [(d.strftime("%d/%m/%Y"), n) for d, n in dados["por_dia"]],
         "Total", total_usa)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
