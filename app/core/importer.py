"""ImportService (§38.17, §39.16) — importação de CSV/Excel com relatório.

O fluxo é sempre em duas etapas, para nunca gravar um lote pela metade:

    1. `ler_tabela()`  -> linhas cruas do arquivo
    2. `validar()`     -> separa válidas de inválidas, com o motivo de cada erro

A gravação só acontece depois que você (ou o usuário, na tela) confirma.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Callable


@dataclass
class LinhaInvalida:
    numero: int
    erros: list[str]
    dados: dict


@dataclass
class ResultadoImportacao:
    validas: list[dict] = field(default_factory=list)
    invalidas: list[LinhaInvalida] = field(default_factory=list)
    colunas_ignoradas: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return len(self.validas) + len(self.invalidas)

    @property
    def ok(self) -> bool:
        return bool(self.validas) and not self.invalidas


def _normalizar(nome: str) -> str:
    import re
    import unicodedata

    valor = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode()
    return re.sub(r"[^\w]+", "_", valor).strip("_").lower()


def ler_tabela(conteudo: bytes, nome_arquivo: str) -> list[dict]:
    """Lê CSV (`;` ou `,`) ou XLSX e devolve as linhas como dicionários."""
    if nome_arquivo.lower().endswith((".xlsx", ".xlsm")):
        return _ler_xlsx(conteudo)
    return _ler_csv(conteudo)


def _ler_csv(conteudo: bytes) -> list[dict]:
    texto = conteudo.decode("utf-8-sig", errors="replace")
    amostra = texto[:2048]
    delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","
    leitor = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    linhas = []
    for bruta in leitor:
        linhas.append({_normalizar(k): (v or "").strip() for k, v in bruta.items() if k})
    return linhas


def _ler_xlsx(conteudo: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.active
    linhas_iter = ws.iter_rows(values_only=True)
    try:
        cabecalho = [_normalizar(c) for c in next(linhas_iter)]
    except StopIteration:
        return []
    linhas = []
    for valores in linhas_iter:
        if all(v is None or str(v).strip() == "" for v in valores):
            continue
        registro = {}
        for chave, valor in zip(cabecalho, valores):
            if not chave:
                continue
            registro[chave] = "" if valor is None else str(valor).strip()
        linhas.append(registro)
    return linhas


def validar(
    linhas: list[dict],
    obrigatorias: list[str],
    conhecidas: list[str] | None = None,
    validador: Callable[[dict], list[str]] | None = None,
) -> ResultadoImportacao:
    """Separa linhas válidas das inválidas.

    `obrigatorias` — colunas que não podem vir vazias.
    `conhecidas`   — colunas aceitas; as demais são reportadas como ignoradas.
    `validador`    — função opcional que recebe a linha e devolve erros extras
                     (ex.: CPF inválido, e-mail duplicado).
    """
    resultado = ResultadoImportacao()
    if conhecidas and linhas:
        extras = sorted(set(linhas[0]) - set(conhecidas))
        resultado.colunas_ignoradas = extras

    for indice, linha in enumerate(linhas, start=2):  # 1 é o cabeçalho
        erros = [f"'{campo}' é obrigatório" for campo in obrigatorias
                 if not linha.get(campo)]
        if not erros and validador is not None:
            erros = validador(linha) or []
        if erros:
            resultado.invalidas.append(LinhaInvalida(indice, erros, linha))
        else:
            resultado.validas.append(linha)
    return resultado


def modelo_csv(colunas: list[str]) -> bytes:
    """Gera um CSV modelo (só o cabeçalho) para o usuário preencher."""
    buffer = io.StringIO()
    csv.writer(buffer, delimiter=";").writerow(colunas)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")
