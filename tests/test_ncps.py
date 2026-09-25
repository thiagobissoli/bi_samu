"""Testes do módulo NCPS."""

import re
from io import BytesIO
from types import SimpleNamespace

from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.database import SessionLocal
from app.core.seeds import ADMIN_EMAIL, ADMIN_SENHA
from app.main import app
from app.modules.ncps import constants as cat
from app.modules.ncps import service
from app.modules.ncps.models import Ncps
from app.modules.ncps.permissions import (filtrar_visiveis, pode_tratar,
                                          pode_triar, pode_ver)

client = TestClient(app)


def _login():
    client.post("/login", data={"email": ADMIN_EMAIL, "senha": ADMIN_SENHA})


def _publicar(**campos) -> tuple[int, str]:
    """Envia o formulário público e devolve (protocolo, código)."""
    dados = {"natureza": "paciente", "descricao": "Paciente caiu da maca.",
             **campos}
    html = client.post("/ncps/publico", data=dados).text
    protocolo = re.search(r'Protocolo</small>\s*<span[^>]*>(\d+)<', html)
    codigo = re.search(r'font-monospace[^>]*>([A-Z0-9]{8})<', html)
    assert protocolo and codigo, "confirmação sem protocolo/código"
    return int(protocolo.group(1)), codigo.group(1)


def _usuario(id_, *permissoes, setores=()):
    return SimpleNamespace(id=id_, empresa_id=1, permissoes=set(permissoes),
                           setores_ncps=set(setores))


# ------------------------------------------------------------------ regras puras

def test_matriz_de_risco_e_a_do_for_samu_038():
    """Mesma escala e faixas da matriz gerada pela IA na Investigação."""
    from app.modules.investigacao.constants import nivel_de_risco

    assert service.cat.nivel_risco(5, 16)["rotulo"] == "Extremo"   # 80
    assert cat.nivel_risco(3, 4)["rotulo"] == "Elevado"            # 12
    assert cat.nivel_risco(2, 2)["rotulo"] == "Moderado"           # 4
    assert cat.nivel_risco(1, 2)["rotulo"] == "Baixo"              # 2
    for p in cat.PROBABILIDADE:
        for c in cat.CONSEQUENCIA:
            assert cat.nivel_risco(p, c)["rotulo"] == nivel_de_risco(p * c)[0]
    assert cat.nivel_risco(2, 3) is None          # 3 não é consequência válida
    assert cat.nivel_risco(None, 4) is None
    assert cat.STATUS["1"] == "Analisando evento"


def test_regras_de_visibilidade():
    paciente = SimpleNamespace(natureza="paciente", confidencial=False,
                               setor_id=None, coordenador_id=None)
    trabalhador = SimpleNamespace(natureza="trabalhador", confidencial=False,
                                  setor_id=3, coordenador_id=None)
    sigilosa = SimpleNamespace(natureza="trabalhador", confidencial=True,
                               setor_id=3, coordenador_id=None)
    qualidade = _usuario(1, "ncps.listar", "ncps.triar_paciente")
    sesmt = _usuario(2, "ncps.listar", "ncps.triar_trabalhador")
    comissao = _usuario(3, "ncps.listar", "ncps.sigilosas")
    coordenador = _usuario(7, "ncps.listar", "ncps.coordenar", setores={3})
    outro_coord = _usuario(8, "ncps.listar", "ncps.coordenar", setores={4})

    assert pode_triar(paciente, qualidade) and not pode_ver(trabalhador, qualidade)
    assert pode_triar(trabalhador, sesmt) and not pode_ver(paciente, sesmt)
    # sigilosa: só a Comissão, nem a Qualidade nem o SESMT
    assert pode_ver(sigilosa, comissao) and pode_triar(sigilosa, comissao)
    assert not pode_ver(sigilosa, coordenador)       # setor não abre sigilosa
    assert not pode_ver(sigilosa, sesmt) and not pode_ver(sigilosa, qualidade)
    # analista vê as do paciente e só as do trabalhador encaminhadas ao seu setor
    assert pode_ver(paciente, coordenador) and not pode_tratar(paciente, coordenador)
    assert pode_ver(trabalhador, coordenador) and pode_tratar(trabalhador, coordenador)
    assert not pode_ver(trabalhador, outro_coord)


def test_consulta_filtra_sigilosas():
    protocolo, _codigo = _publicar(natureza="trabalhador",
                                   tipo_evento_trab="violencia_assedio",
                                   descricao="Assédio no plantão.")
    db = SessionLocal()
    try:
        def ids(usuario):
            q = filtrar_visiveis(select(Ncps.id).where(Ncps.id == protocolo), usuario)
            return list(db.scalars(q))
        assert ids(_usuario(99, "ncps.triar_trabalhador")) == []
        assert ids(_usuario(99, "ncps.sigilosas")) == [protocolo]
    finally:
        db.close()


# ------------------------------------------------------------------ páginas públicas

def test_formulario_publico_e_acompanhamento():
    protocolo, codigo = _publicar()
    db = SessionLocal()
    try:
        n = db.get(Ncps, protocolo)
        assert n.anonima and n.notificante_id is None
        assert n.codigo_hash and codigo not in (n.codigo_hash or "")  # só o hash
    finally:
        db.close()

    ok = client.post("/ncps/acompanhar",
                     data={"protocolo": protocolo, "codigo": codigo.lower()}).text
    assert f"Protocolo {protocolo}" in ok and "Recebida" in ok
    assert "caiu da maca" not in ok            # o relato não é exposto

    errado = client.post("/ncps/acompanhar",
                         data={"protocolo": protocolo, "codigo": "XXXXXXXX"}).text
    assert "não encontrada" in errado


def test_formulario_publico_valida_campos():
    html = client.post("/ncps/publico",
                       data={"natureza": "trabalhador", "descricao": ""}).text
    assert "Descreva o que aconteceu" in html


def test_telas_internas_exigem_login():
    novo = TestClient(app)
    for url in ("/ncps/", "/ncps/notificar", "/ncps/painel"):
        resp = novo.get(url, headers={"accept": "text/html"}, follow_redirects=False)
        assert resp.status_code == 303 and resp.headers["location"] == "/login"


# ------------------------------------------------------------------ tratativa

def test_fluxo_completo_de_tratativa():
    _login()
    client.post("/ncps/cadastros/gestores", data={"nome": "Gestor Teste NCPS"})
    client.post("/ncps/cadastros/locais", data={"nome": "Base Teste NCPS"})

    html = client.post("/ncps/notificar", data={
        "natureza": "trabalhador", "tipo_evento_trab": "acidente_tipico",
        "descricao": "Corte com perfurocortante.", "houve_lesao": "sim",
        "data_hora_ocorrencia": "2026-09-20T14:30"}).text
    assert "Notificação registrada" in html
    ncps_id = int(re.search(r'Protocolo</small>\s*<span[^>]*>(\d+)<', html).group(1))

    db = SessionLocal()
    try:
        n = db.get(Ncps, ncps_id)
        # acidente exige revisão do PGR (NR-1 1.5.4.4.6 "d")
        assert n.ocupacional and n.ocupacional.revisar_pgr
        admin_id = n.notificante_id
        from app.modules.ncps.models import NcpsGestor, NcpsLocal
        gestor = db.scalar(select(NcpsGestor).where(NcpsGestor.nome == "Gestor Teste NCPS"))
        local = db.scalar(select(NcpsLocal).where(NcpsLocal.nome == "Base Teste NCPS"))
    finally:
        db.close()

    assert client.get(f"/ncps/{ncps_id}").status_code == 200

    def salvar(**dados):
        resp = client.post(f"/ncps/{ncps_id}", data=dados, follow_redirects=False)
        assert resp.status_code == 303, resp.text
        assert "erro=" not in resp.headers["location"], resp.headers["location"]

    client.post("/ncps/setores",
                data={"nome": "Setor Teste NCPS", "usuarios": [admin_id]})
    db = SessionLocal()
    try:
        from app.modules.ncps.models import NcpsSetor
        setor = db.scalar(select(NcpsSetor).where(NcpsSetor.nome == "Setor Teste NCPS"))
        assert [u.id for u in setor.usuarios] == [admin_id]
    finally:
        db.close()
    salvar(secao="triagem", natureza="trabalhador", procedente="2", status="1",
           local_id=local.id, gestor_id=gestor.id, setor_id=setor.id)
    salvar(secao="classificacao", tipo_evento_trab="acidente_tipico",
           afastamento="sim", dias_afastamento="3", cat_emitida="sim",
           cat_numero="123", cat_data="2026-09-21")
    salvar(secao="risco", inicial_ps="3-4", inicial_justificativa="Recorrente",
           residual_ps="2-3")        # consequência 3 não existe: ignorada
    salvar(secao="analise", cronologia="14h30 corte", problemas="Descarte")
    salvar(secao="causa_add", metodo="ishikawa", categoria="metodo",
           descricao="Sem caixa de descarte", causa_raiz="1")
    salvar(secao="acao_add", descricao="Instalar caixas de descarte",
           tipo="epc", responsavel="SESMT", prazo="2020-01-01")

    db = SessionLocal()
    try:
        n = db.get(Ncps, ncps_id)
        assert n.status == "1" and n.setor_id == setor.id
        # os analistas do setor são avisados e o alerta aparece no Início
        from app.models import Notificacao, Usuario
        assert db.scalar(select(Notificacao).where(
            Notificacao.usuario_id == admin_id,
            Notificacao.titulo == f"NCPS #{ncps_id} encaminhada ao setor Setor Teste NCPS"))
        from app.modules.inicio.service import alertas
        admin = db.get(Usuario, admin_id)
        assert any("aguardando análise do seu setor" in a["titulo"]
                   for a in alertas(db, admin))
        assert n.gestor_id == gestor.id and n.local_id == local.id
        assert n.ocupacional.dias_afastamento == 3 and n.ocupacional.cat_numero == "123"
        assert n.risco("inicial").nivel["rotulo"] == "Elevado"   # 3 × 4 = 12
        assert n.risco("residual") is None
        assert n.analise.cronologia == "14h30 corte"
        assert n.causas[0].causa_raiz
        assert service.acoes_atrasadas(n)                      # prazo vencido
        acao_id = n.acoes[0].id
    finally:
        db.close()

    salvar(secao="acao_update", acao_id=acao_id, status="concluida",
           eficacia="eficaz")
    db = SessionLocal()
    try:
        acao = db.get(Ncps, ncps_id).acoes[0]
        assert acao.status == "concluida" and acao.concluida_em
    finally:
        db.close()

    # indicadores, lista e exportação enxergam a notificação
    dados = client.get("/ncps/api/indicadores").json()["data"]
    assert dados["trabalhador"] >= 1 and dados["ocupacional"]["cats"] >= 1
    assert client.get("/ncps/painel").status_code == 200
    assert f'/ncps/{ncps_id}"' in client.get("/ncps/?q=perfurocortante").text

    planilha = client.get("/ncps/exportar?q=perfurocortante")
    assert planilha.status_code == 200
    from openpyxl import load_workbook
    aba = load_workbook(BytesIO(planilha.content)).active
    cabecalho = [c.value for c in aba[1]]
    assert "Protocolo" in cabecalho and "Risco antes" in cabecalho and "Setor" in cabecalho
    assert any(row[0] == ncps_id for row in aba.iter_rows(min_row=2, values_only=True))

    # exclusão lógica
    client.post(f"/ncps/{ncps_id}/excluir")
    assert client.get(f"/ncps/{ncps_id}").status_code == 404


def test_secao_invalida_nao_quebra():
    _login()
    protocolo, _ = _publicar()
    resp = client.post(f"/ncps/{protocolo}", data={"secao": "causa_add",
                                                  "metodo": "londres"},
                       follow_redirects=False)
    assert resp.status_code == 303 and "erro=" in resp.headers["location"]


def test_powerbi_exige_token():
    assert client.get("/ncps/api/powerbi").status_code == 401
    assert client.get("/ncps/api/powerbi",
                      headers={"Authorization": "Bearer errado"}).status_code == 401


def test_pgr_carrega_catalogo_inicial():
    _login()
    html = client.get("/ncps/pgr").text
    assert "Radioperação / TARM" in html and "Agentes biológicos" in html


def test_serie_mensal_sem_buracos():
    """Meses sem notificação entram zerados: o eixo não pode pular meses."""
    _login()
    meses = [m["mes"] for m in
             client.get("/ncps/api/indicadores").json()["data"]["por_mes"]]
    assert len(meses) == 12
    for anterior, atual in zip(meses, meses[1:]):
        a, m = map(int, anterior.split("-"))
        assert atual == (f"{a:04d}-{m + 1:02d}" if m < 12 else f"{a + 1:04d}-01")
