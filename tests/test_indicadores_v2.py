"""Indicadores: filtros ampliados, carga sob demanda, detalhamento e correlação."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.core.seeds import ADMIN_EMAIL, ADMIN_SENHA
from app.main import app
from app.modules.indicadores import correlacao as C
from app.modules.indicadores import filtros as F
from app.modules.indicadores import nucleo
from app.modules.indicadores.service import IndicadoresService

client = TestClient(app)
HTML = {"accept": "text/html"}


def _login():
    client.post("/login", data={"email": ADMIN_EMAIL, "senha": ADMIN_SENHA})


@pytest.fixture(scope="module")
def periodo():
    """Últimos 30 dias da base, como a tela abre."""
    return F.padrao(IndicadoresService(1).periodo_disponivel_leve())


def _qs(f: dict, **extra) -> str:
    return F.query_string(f, **extra)


# ------------------------------------------------------------ filtros

def test_padrao_sao_os_ultimos_30_dias():
    f = F.padrao("2026-09-25")
    assert (f["data_inicial"], f["data_final"]) == ("2026-08-27", "2026-09-25")


def test_filtro_de_hora_vira_a_meia_noite(periodo):
    df = F.aplicar(nucleo.carregar(1), periodo)
    noite = F.aplicar(df, {"hora_min": "19", "hora_max": "6"})
    assert set(noite["hora"].unique()) <= set(range(19, 24)) | set(range(0, 7))
    manha = F.aplicar(df, {"hora_min": "8", "hora_max": "10"})
    assert set(manha["hora"].unique()) <= {8, 9, 10}


def test_filtro_por_faixa_de_um_indicador_de_tempo(periodo):
    df = F.aplicar(nucleo.carregar(1), periodo)
    sub = F.aplicar(df, {"tempo_metrica": "tempo_resposta",
                         "tempo_min": "10", "tempo_max": "15"})
    assert len(sub) and sub["tempo_resposta"].between(600, 900).all()


def test_filtros_simnao_e_dia_da_semana(periodo):
    df = F.aplicar(nucleo.carregar(1), periodo)
    fora = F.aplicar(df, {"fora_municipio": "sim"})
    assert len(fora) and fora["fora_do_municipio"].eq(True).all()
    domingo = F.aplicar(df, {"dia_semana": ["Domingo"]})
    assert set(domingo["plantao_dia_semana"].unique()) == {"Domingo"}


def test_nome_antigo_do_filtro_iscm_continua_valendo():
    assert F.ler({"iscmv": "1"})["iscm"] == "1"


# ------------------------------------------------------------ carga sob demanda

def test_dashboard_so_calcula_depois_de_aplicar(periodo):
    _login()
    vazio = client.get("/indicadores/tempo-resposta", headers=HTML).text
    assert "clique em <strong>Aplicar</strong>" in vazio
    assert "chart-1" not in vazio
    assert f'value="{periodo["data_inicial"]}"' in vazio   # 30 dias preenchidos

    cheio = client.get(f"/indicadores/tempo-resposta?{_qs(periodo)}&aplicar=1",
                       headers=HTML).text
    assert "chart-1" in cheio and "modal-ocorrencias" in cheio
    assert "mapa-calor" in cheio or '"heatmap"' in cheio


def test_opcoes_dos_filtros_vem_a_parte():
    _login()
    dados = client.get("/indicadores/api/opcoes").json()["data"]
    assert "dia_semana" in dados and "regulador" in dados
    assert ["vermelho", "Vermelho"] in dados["cor"]


# ------------------------------------------------------------ detalhamento

def test_detalhamento_reproduz_o_valor_do_grafico(periodo):
    _login()
    d = client.get(f"/indicadores/api/tempo-central?{_qs(periodo)}").json()["data"]
    i = next(k for k, c in enumerate(d["charts"])
             if c["titulo"].startswith("Tempo de Central por hora do dia"))
    x = 10
    valor = d["charts"][i]["datasets"][0]["data"][x]
    r = client.get(f"/indicadores/api/tempo-central/ocorrencias?{_qs(periodo)}"
                   f"&grafico={i}&x={x}&s=0").json()["data"]
    assert r["total"] > 0 and "Tempo de Central (mm:ss)" in r["colunas"]
    # a média dos tempos listados é o valor do ponto
    col = r["colunas"].index("Tempo de Central (mm:ss)")
    segundos = []
    for pagina in range(1, r["paginas"] + 1):
        p = client.get(f"/indicadores/api/tempo-central/ocorrencias?{_qs(periodo)}"
                       f"&grafico={i}&x={x}&s=0&pagina={pagina}").json()["data"]
        for linha in p["linhas"]:
            m, s = linha[col].split(":")
            segundos.append(int(m) * 60 + int(s))
    assert abs(sum(segundos) / len(segundos) / 60 - valor) < 0.05


def test_contagem_do_detalhamento_bate_com_a_barra(periodo):
    _login()
    d = client.get(f"/indicadores/api/codigos?{_qs(periodo)}").json()["data"]
    barra = d["charts"][0]
    r = client.get(f"/indicadores/api/codigos/ocorrencias?{_qs(periodo)}"
                   "&grafico=0&x=0&s=0").json()["data"]
    assert r["total"] == barra["datasets"][0]["data"][0]
    assert r["recorte"].startswith(str(barra["labels"][0]))


def test_todos_os_graficos_dos_dashboards_sao_clicaveis(periodo):
    from app.modules.indicadores.constants import TEMAS

    svc = IndicadoresService(1)
    sem = [(t, c["titulo"]) for t in TEMAS for c in svc.dashboard(t, periodo)["charts"]
           if not c["drill"]]
    assert sem == []


def test_exporta_ocorrencias_em_excel(periodo):
    from openpyxl import load_workbook

    _login()
    resp = client.get(f"/indicadores/codigos/ocorrencias.xlsx?{_qs(periodo)}"
                      "&grafico=0&x=0&s=0")
    assert resp.status_code == 200
    aba = load_workbook(BytesIO(resp.content)).active
    assert [c.value for c in aba[2]][:2] == ["Ocorrência", "Data/hora"]
    assert aba.max_row > 2


def test_grafico_sem_indice_valido_nao_quebra(periodo):
    _login()
    r = client.get(f"/indicadores/api/codigos/ocorrencias?{_qs(periodo)}&grafico=999&x=0")
    assert r.status_code == 400


# ------------------------------------------------------------ correlação

def test_p_valor_confere_com_a_tabela_t():
    assert abs(C.p_valor(0.5, 20) - 0.0248) < 0.0005
    assert abs(C.p_valor(0.3, 50) - 0.0343) < 0.0005


def test_dispersao_por_dia(periodo):
    _login()
    d = client.get(f"/indicadores/api/correlacao?{_qs(periodo)}&modo=dispersao"
                   "&x=t_central&y=tempo_resposta&dimensao=dia&min_n=5").json()["data"]
    assert d["n_grupos"] >= 10 and -1 <= d["pearson"] <= 1
    assert "Correlação não prova causa" in d["interpretacao"]
    # clicar num ponto lista as ocorrências daquele dia
    grupo = d["pontos"][0]["grupo"]
    r = client.get(f"/indicadores/api/correlacao/ocorrencias?{_qs(periodo)}"
                   f"&dimensao=dia&grupo={grupo}&metrica=tempo_resposta").json()["data"]
    assert r["total"] > 0


def test_matriz_e_series(periodo):
    _login()
    m = client.get(f"/indicadores/api/correlacao?{_qs(periodo)}&modo=matriz"
                   "&metricas=t_central&metricas=tempo_resposta&metricas=empenhos"
                   "&dimensao=unidade").json()["data"]
    assert len(m["celulas"]) == 3 and m["celulas"][0][0]["r"] == 1.0
    assert m["celulas"][0][1]["r"] == m["celulas"][1][0]["r"]          # simétrica

    s = client.get(f"/indicadores/api/correlacao?{_qs(periodo)}&modo=series"
                   "&metricas=tempo_resposta&metricas=desperdicio&dimensao=semana"
                   ).json()["data"]
    assert len(s["series"]) == 2 and len(s["grupos"]) >= 3
    validos = [v for v in s["indexadas"][0]["valores"] if v is not None]
    assert abs(sum(validos) / len(validos) - 100) < 1       # base 100 = média


def test_pagina_de_correlacao(periodo):
    _login()
    assert "clique em <strong>Calcular</strong>" in client.get(
        "/indicadores/correlacao", headers=HTML).text
    html = client.get(f"/indicadores/correlacao?{_qs(periodo)}&aplicar=1&modo=matriz"
                      "&dimensao=unidade", headers=HTML).text
    assert "matriz-corr" in html
