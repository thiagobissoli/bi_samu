"""Documento "Saída de USA" calculado dos dados do vSky que o sistema importa."""

import io
from datetime import date, datetime

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.core.seeds import ADMIN_EMAIL, ADMIN_SENHA
from app.main import app
from app.modules.sesa import prazos, saida_usa as SU, service

client = TestClient(app)
HTML = {"accept": "text/html"}


def _login():
    client.post("/login", data={"email": ADMIN_EMAIL, "senha": ADMIN_SENHA})


def _df(linhas: list[dict]) -> pd.DataFrame:
    """DataFrame no formato do núcleo, só com as colunas que a Saída de USA usa."""
    base = {"dt_ocorr": pd.NaT, "dia": None, "unidade_curta": None,
            "dt_inicio_deslocamento": pd.NaT, "recurso": None, "codigo_cor": None,
            "cidade": None}
    return pd.DataFrame([{**base, **l} for l in linhas])


def _linha(dia, unidade, recurso, cor, cidade, saiu=True):
    d = datetime(dia.year, dia.month, dia.day, 8, 0)
    return {"dt_ocorr": d, "dia": dia, "unidade_curta": unidade, "recurso": recurso,
            "codigo_cor": cor, "cidade": cidade,
            "dt_inicio_deslocamento": d if saiu else pd.NaT}


# ------------------------------------------------------------------ cálculo puro

def test_so_conta_empenhos_que_sairam():
    df = _df([
        _linha(date(2026, 9, 1), "USA 10", "USA", "vermelho", "VITORIA"),
        _linha(date(2026, 9, 1), "USA 10", "USA", "amarelo", "VITORIA"),
        _linha(date(2026, 9, 2), "USB 42", "USB", "verde", "SERRA"),
        _linha(date(2026, 9, 2), "USA 20", "USA", "vermelho", "SERRA", saiu=False),  # não saiu
    ])
    d = SU.calcular(df, date(2026, 9, 1), date(2026, 9, 30))
    assert d["total"] == 3 and d["usa"] == 2 and d["usb"] == 1
    assert d["dias"] == 2
    assert d["por_unidade"] == [("USA 10", 2)]
    assert dict(d["por_cor"]) == {"Vermelho": 1, "Amarelo": 1}
    assert dict(d["por_municipio"]) == {"VITORIA": 2}


def test_filtra_pelo_periodo():
    df = _df([
        _linha(date(2026, 8, 31), "USA 10", "USA", "vermelho", "VITORIA"),   # fora
        _linha(date(2026, 9, 15), "USA 10", "USA", "amarelo", "VITORIA"),    # dentro
        _linha(date(2026, 10, 1), "USA 10", "USA", "verde", "VITORIA"),      # fora
    ])
    d = SU.calcular(df, date(2026, 9, 1), date(2026, 9, 30))
    assert d["usa"] == 1 and d["por_dia"] == [(date(2026, 9, 15), 1)]


def test_planilha_tem_as_abas():
    from openpyxl import load_workbook

    df = _df([_linha(date(2026, 9, 1), "USA 10", "USA", "vermelho", "VITORIA"),
              _linha(date(2026, 9, 2), "USB 42", "USB", "verde", "SERRA")])
    d = SU.calcular(df, date(2026, 9, 1), date(2026, 9, 30))
    wb = load_workbook(io.BytesIO(SU.gerar_xlsx(d, "setembro/2026")))
    assert wb.sheetnames == ["Resumo", "Por unidade", "Por município",
                             "Por código", "Por dia"]
    resumo = wb["Resumo"]
    assert resumo["A1"].value.startswith("Saída de USA")
    # a linha "Saídas de USA" tem o valor 1
    valores = {resumo.cell(r, 1).value: resumo.cell(r, 2).value for r in range(5, 10)}
    assert valores["Saídas de USA"] == 1


def test_periodo_da_competencia():
    assert prazos.periodo_competencia_anterior(date(2026, 10, 1)) == \
        (date(2026, 9, 1), date(2026, 9, 30))


# ------------------------------------------------------------------ fluxo (dados reais)

def test_pagina_e_geracao(monkeypatch, tmp_path):
    from app.core.config import settings
    from app.core.database import SessionLocal
    from app.modules.indicadores import nucleo

    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    comp = "2026-10"                     # período = setembro/2026 (há dados na cópia)
    _login()

    # a página abre com o resumo e o botão
    html = client.get(f"/sesa/saida-usa/{comp}?para=convenio_007", headers=HTML).text
    assert "Saída de USA" in html and "Gerar planilha e anexar" in html

    # confere que o item é mesmo o "Saída de USA" do convênio 007
    db = SessionLocal()
    try:
        obrig, item = service.item_do_gerador(db, 1, "convenio_007", "/sesa/saida-usa")
        assert item is not None and "USA" in item.nome
    finally:
        db.close()

    r = client.post(f"/sesa/saida-usa/{comp}/gerar", data={"para": "convenio_007"},
                    follow_redirects=False)
    assert r.headers["location"].startswith(f"/sesa/convenio_007/{comp}")

    db = SessionLocal()
    try:
        o = service.obrigacao_por_chave(db, 1, "convenio_007")
        dados = service.detalhe(db, 1, o, prazos.competencia_de(comp), date(2026, 10, 5))
        linha = next(l for l in dados["linhas"] if l["item"].chave == "i04")
        assert linha["gerador"].startswith("/sesa/saida-usa/") and "para=convenio_007" in linha["gerador"]
        anexo = linha["anexo"]
        assert anexo is not None
        assert anexo.arquivo.nome_original == "Saida de USA - 09-2026.xlsx"
        # o xlsx abre e o número de USA bate com o cálculo direto
        from openpyxl import load_workbook

        from app.core.storage import absolute_path
        wb = load_workbook(absolute_path(anexo.arquivo))
        resumo = wb["Resumo"]
        valor_usa = {resumo.cell(r, 1).value: resumo.cell(r, 2).value
                     for r in range(5, 10)}["Saídas de USA"]
        esperado = SU.calcular(nucleo.carregar(1), date(2026, 9, 1), date(2026, 9, 30))["usa"]
        assert valor_usa == esperado > 0
    finally:
        db.close()
